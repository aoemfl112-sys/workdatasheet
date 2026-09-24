"""우측 표 문구. 사람이 적은 내용만 지정한 칸에 넣는다.

비어 있으면 템플릿을 그대로 둔다. AI가 문단을 만들지 않는다.
칸 위치는 레이아웃 JSON의 text_fields 로 정의한다.
"""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

ALLOWED_WIDGETS = {"text_input", "text_area"}


def fields_from_layout(layout: dict) -> list[dict]:
    items = []
    for raw in layout.get("text_fields") or []:
        field_id = str(raw.get("id") or "").strip()
        label = str(raw.get("label") or "").strip()
        if not field_id or not label:
            continue
        widget = str(raw.get("widget") or "text_input")
        if widget not in ALLOWED_WIDGETS:
            widget = "text_input"
        items.append(
            {
                "id": field_id,
                "label": label,
                "widget": widget,
                "placeholder": str(raw.get("placeholder") or ""),
                "font_pt": int(raw.get("font_pt") or 10),
                "mode": str(raw.get("mode") or "cell"),
                "targets": list(raw.get("targets") or []),
            }
        )
    return items


def nonempty_values(values: dict[str, str] | None) -> dict[str, str]:
    result: dict[str, str] = {}
    for key, raw in (values or {}).items():
        text = (raw or "").strip()
        if text:
            result[str(key)] = text
    return result


def _table_by_names(slide, names: list[str]):
    wanted = {name for name in names if name}
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.TABLE:
            continue
        if shape.name in wanted:
            return shape
    return None


def _set_cell_text(cell, text: str, font_pt: int) -> None:
    lines = text.replace("\r\n", "\n").split("\n")
    if not lines:
        lines = [""]
    tf = cell.text_frame
    tf.word_wrap = True
    while len(tf.paragraphs) < len(lines):
        tf.add_paragraph()
    for index, paragraph in enumerate(tf.paragraphs):
        paragraph.text = lines[index] if index < len(lines) else ""
        paragraph.font.size = Pt(font_pt)
        paragraph.font.color.rgb = RGBColor(40, 40, 40)


def _write_cell(table, row: int, col: int, text: str, font_pt: int) -> bool:
    if row < 0 or col < 0:
        return False
    if row >= len(table.rows) or col >= len(table.columns):
        return False
    _set_cell_text(table.cell(row, col), text, font_pt)
    return True


def apply_text_fields(slide, layout: dict, values: dict[str, str] | None) -> list[str]:
    """입력된 필드만 표에 기입한다. 적용한 필드 id 목록을 돌려준다."""
    filled: list[str] = []
    written = nonempty_values(values)
    if not written:
        return filled
    for field in fields_from_layout(layout):
        text = written.get(field["id"])
        if not text:
            continue
        applied = False
        for target in field["targets"]:
            names = [str(target.get("shape_name") or "")]
            names.extend(str(n) for n in (target.get("alt_names") or []) if n)
            shape = _table_by_names(slide, names)
            if shape is None:
                continue
            table = shape.table
            mode = str(target.get("mode") or field["mode"])
            font_pt = int(target.get("font_pt") or field["font_pt"])
            if mode == "rows":
                col = int(target.get("col") or 1)
                row_start = int(target.get("row_start") or 0)
                max_rows = int(target.get("max_rows") or (len(table.rows) - row_start))
                lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
                if not lines:
                    continue
                last_index = min(max_rows, len(table.rows) - row_start) - 1
                if last_index < 0:
                    continue
                for offset, line in enumerate(lines):
                    if offset < last_index:
                        if _write_cell(table, row_start + offset, col, line, font_pt):
                            applied = True
                    else:
                        leftover = "\n".join(lines[offset:])
                        if _write_cell(table, row_start + last_index, col, leftover, font_pt):
                            applied = True
                        break
            else:
                row = int(target.get("row") or 0)
                col = int(target.get("col") or 0)
                if _write_cell(table, row, col, text, font_pt):
                    applied = True
        if applied:
            filled.append(field["id"])
    return filled
