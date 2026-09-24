"""템플릿 복사 후 슬롯에 사진·라벨을 삽입한다."""

from __future__ import annotations

import copy
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

from engine.image_ops import centered_origin, contain_size, image_pixel_size, prepare_image
from engine.layout import load_layout, slot_by_id
from engine.slot_edit import photo_area
from engine.text_fields import apply_text_fields

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
EMBED_ATTR = f"{{{R_NS}}}embed"
LINK_ATTR = f"{{{R_NS}}}link"
RID_ATTR = f"{{{R_NS}}}id"
EMU_PER_CM = 360000
SAMPLE_BADGES = {"1", "2", "3", "4"}


@dataclass
class SlotAssignment:
    slot_id: int
    image_path: str | Path
    compress: bool = False
    label: str = ""


@dataclass
class SlidePage:
    assignments: list[SlotAssignment]
    text_values: dict[str, str] = field(default_factory=dict)
    layout_data: dict | None = None


def _shape_plain_text(shape) -> str:
    try:
        if not getattr(shape, "has_text_frame", False):
            return ""
        return " ".join(
            part.text.strip()
            for part in shape.text_frame.paragraphs
            if part.text.strip()
        ).strip()
    except Exception:
        return ""


def _in_work_drawing(shape, layout: dict) -> bool:
    """우측 표가 아닌 작업 도면 쪽에 있는 도형인지 본다."""
    area = photo_area(layout)
    x = float(shape.left) / EMU_PER_CM
    return x < area["x_cm"] + area["width_cm"] + 0.15


def _is_sample_drawing_chrome(shape, layout: dict) -> bool:
    """예제에서 남은 번호 칸·안내 문구. 사진·사용자가 적은 라벨은 건드리지 않는다."""
    name = (getattr(shape, "name", "") or "").strip().upper()
    if name.startswith(("PHOTO_", "SLOT_", "LABEL_")):
        return False
    if getattr(shape, "has_table", False):
        return False
    if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER:
        return False
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return False
    if not _in_work_drawing(shape, layout):
        return False
    text = _shape_plain_text(shape)
    if not text:
        return False
    if text in SAMPLE_BADGES:
        return True
    upper = text.upper()
    if "TEST PROGRAM" in upper:
        return True
    if "운전자표출기" in text:
        return True
    return False


def _shape_map(slide) -> dict[str, object]:
    return {shape.name: shape for shape in slide.shapes}


def _remove_shape(shape) -> None:
    element = shape._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _remove_drawing_samples(slide, layout: dict) -> None:
    """작업 도면의 예제 번호·안내 문구를 지운다. 없는 상태가 기본값이다."""
    for shape in list(slide.shapes):
        if _is_sample_drawing_chrome(shape, layout):
            _remove_shape(shape)


def _iter_shape_holders(prs: Presentation):
    for master in prs.slide_masters:
        yield master
        for layout in master.slide_layouts:
            yield layout
    for slide in prs.slides:
        yield slide


def _find_table(shapes, name: str):
    for shape in shapes:
        if getattr(shape, "name", "") == name and getattr(shape, "has_table", False):
            return shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            found = _find_table(shape.shapes, name)
            if found is not None:
                return found
    return None


def _clear_table_cell(table, row: int, col: int) -> bool:
    if row < 0 or col < 0:
        return False
    if row >= len(table.rows) or col >= len(table.columns):
        return False
    cell = table.cell(row, col)
    tf = cell.text_frame
    for paragraph in tf.paragraphs:
        paragraph.text = ""
    return True


def clear_header_fields(prs: Presentation, layout: dict) -> None:
    """MODEL·문서번호·Rev·작성일 등 예제 머리글을 비운다. 라벨 칸은 그대로 둔다."""
    targets = list(layout.get("header_clear") or [])
    if not targets:
        return
    for holder in _iter_shape_holders(prs):
        for target in targets:
            shape = _find_table(holder.shapes, str(target.get("shape_name") or ""))
            if shape is None:
                continue
            _clear_table_cell(
                shape.table,
                int(target.get("row") or 0),
                int(target.get("col") or 0),
            )


def _set_label_text(shape, text: str, font_pt: int) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(font_pt)
    p.font.bold = True
    p.font.color.rgb = RGBColor(40, 40, 40)
    for extra in tf.paragraphs[1:]:
        extra.text = ""


def _add_label_box(slide, slot: dict, text: str, font_pt: int):
    box = slide.shapes.add_textbox(
        Cm(slot["label_x_cm"]),
        Cm(slot["label_y_cm"]),
        Cm(slot["label_width_cm"]),
        Cm(slot["label_height_cm"]),
    )
    box.name = f"LABEL_{slot['id']}"
    _set_label_text(box, text, font_pt)
    return box


def _remap_rids(element, rid_map: dict[str, str]) -> None:
    for node in element.iter():
        for attr in (EMBED_ATTR, LINK_ATTR, RID_ATTR):
            value = node.attrib.get(attr)
            if value in rid_map:
                node.attrib[attr] = rid_map[value]


def duplicate_slide(prs: Presentation, index: int):
    """빈 템플릿 슬라이드를 복제한다. 이미지 관계 rId 도 함께 맞춘다."""
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)
    dest_tree = dest.shapes._spTree
    for child in list(dest_tree):
        tag = child.tag
        if tag in {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}:
            continue
        dest_tree.remove(child)

    rid_map: dict[str, str] = {}
    for r_id, rel in source.part.rels.items():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        if rel.is_external:
            new_id = dest.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            new_id = dest.part.relate_to(rel.target_part, rel.reltype)
        rid_map[r_id] = new_id

    source_tree = source.shapes._spTree
    for child in source_tree:
        tag = child.tag
        if tag in {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}:
            continue
        cloned = copy.deepcopy(child)
        _remap_rids(cloned, rid_map)
        dest_tree.append(cloned)
    return dest


def _parse_slot_name(name: str) -> int | None:
    text = (name or "").strip().upper()
    if text.startswith("SLOT_") and text[5:].isdigit():
        return int(text[5:])
    if text.startswith("LABEL_") and text[6:].isdigit():
        return int(text[6:])
    return None


def _set_shape_box(shape, x_cm: float, y_cm: float, width_cm: float, height_cm: float) -> None:
    shape.left = Cm(x_cm)
    shape.top = Cm(y_cm)
    shape.width = Cm(width_cm)
    shape.height = Cm(height_cm)


def _sync_slot_placeholders(slide, layout: dict) -> None:
    """빈 칸 박스(SLOT_*)는 지우고, 라벨 칸만 현재 좌표에 맞춘다."""
    live_ids = {int(slot["id"]) for slot in layout.get("slots") or []}
    for shape in list(slide.shapes):
        name = (getattr(shape, "name", "") or "").strip()
        upper = name.upper()
        slot_id = _parse_slot_name(name)
        if upper.startswith("SLOT_"):
            _remove_shape(shape)
            continue
        if slot_id is not None and slot_id not in live_ids:
            _remove_shape(shape)

    shapes = _shape_map(slide)
    for slot in layout.get("slots") or []:
        label_name = f"LABEL_{slot['id']}"
        label_shape = shapes.get(label_name)
        if label_shape is None:
            continue
        _set_shape_box(
            label_shape,
            float(slot["label_x_cm"]),
            float(slot["label_y_cm"]),
            float(slot["label_width_cm"]),
            float(slot["label_height_cm"]),
        )


def _fill_slide(slide, layout: dict, assignments: list[SlotAssignment], text_values: dict[str, str] | None) -> None:
    _remove_drawing_samples(slide, layout)
    _sync_slot_placeholders(slide, layout)
    shapes = _shape_map(slide)
    font_pt = int(layout.get("label_font_pt", 11))
    used_slots: set[int] = set()
    for item in assignments:
        if item.slot_id in used_slots:
            raise ValueError(f"슬롯 {item.slot_id} 에 사진이 두 번 지정되었습니다.")
        used_slots.add(item.slot_id)
        slot = slot_by_id(layout, item.slot_id)
        image_path = Path(item.image_path)
        if not image_path.exists():
            raise FileNotFoundError(image_path)

        px_w, px_h = image_pixel_size(image_path)
        dest_w, dest_h = contain_size(px_w, px_h, slot["width_cm"], slot["height_cm"])
        left, top = centered_origin(
            slot["x_cm"],
            slot["y_cm"],
            slot["width_cm"],
            slot["height_cm"],
            dest_w,
            dest_h,
        )
        stream = prepare_image(image_path, dest_w, dest_h, compress=item.compress)
        picture = slide.shapes.add_picture(stream, Cm(left), Cm(top), Cm(dest_w), Cm(dest_h))
        picture.name = f"PHOTO_{slot['id']}"

        label_text = (item.label or "").strip()
        label_name = f"LABEL_{slot['id']}"
        existing_label = shapes.get(label_name)
        if label_text:
            if existing_label is not None:
                _set_label_text(existing_label, label_text, font_pt)
            else:
                _add_label_box(slide, slot, label_text, font_pt)
        elif existing_label is not None:
            _remove_shape(existing_label)

    for shape in list(slide.shapes):
        name = (getattr(shape, "name", "") or "").upper()
        slot_id = _parse_slot_name(name)
        if name.startswith("LABEL_") and slot_id not in used_slots:
            _remove_shape(shape)

    apply_text_fields(slide, layout, text_values)


def build_pptx(
    template_path: str | Path,
    layout_path: str | Path,
    assignments: list[SlotAssignment],
    output_path: str | Path,
    text_values: dict[str, str] | None = None,
    pages: list[SlidePage] | None = None,
    layout_data: dict | None = None,
) -> Path:
    """템플릿을 복사한 뒤 지정된 슬롯에 사진과 라벨을 넣고 저장한다."""
    layout = layout_data if layout_data is not None else load_layout(layout_path)
    template_path = Path(template_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    slide_pages = pages
    if slide_pages is None:
        slide_pages = [SlidePage(list(assignments), dict(text_values or {}))]
    usable = [page for page in slide_pages if page.assignments]
    if not usable:
        raise ValueError("배치할 사진이 없습니다. 슬롯을 지정하세요.")

    prs = Presentation(str(output_path))
    clear_header_fields(prs, layout)
    slide_index = int(layout.get("slide_index", 0))
    if slide_index >= len(prs.slides):
        raise IndexError(f"slide_index {slide_index} 가 범위를 벗어났습니다.")

    while len(prs.slides) < slide_index + len(usable):
        duplicate_slide(prs, slide_index)

    for offset, page in enumerate(usable):
        page_layout = page.layout_data if page.layout_data is not None else layout
        _fill_slide(
            prs.slides[slide_index + offset],
            page_layout,
            page.assignments,
            page.text_values,
        )

    prs.save(str(output_path))
    return output_path


def count_pictures(pptx_path: str | Path, slide_index: int = 0) -> int:
    prs = Presentation(str(pptx_path))
    slide = prs.slides[slide_index]
    return sum(1 for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
